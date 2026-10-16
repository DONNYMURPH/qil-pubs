"""Read the cancer center's trial list."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

NCT_RE = re.compile(r"NCT\d{8}")
STRICT_NCT_RE = re.compile(r"^NCT\d{8}$")

# Column headers as the cancer center exports them. Matching is done on a
# normalized (lowercased, de-spaced) form so minor header edits don't break us.
COLUMNS = {
    "trialprotocolnumber": "protocol",
    "trialnctnumber": "nct",
    "trialphase": "phase",
    "trialname": "name",
    "trialtrialstatus": "status",
}


@dataclass
class Trial:
    protocol: str = ""
    nct: str = ""
    phase: str = ""
    name: str = ""
    status: str = ""
    row: int = 0
    #: NCT numbers we filled in ourselves, kept separate from the official cell
    resolved_nct: str = ""
    notes: list[str] = field(default_factory=list)

    @property
    def effective_nct(self) -> str:
        return self.nct or self.resolved_nct


def _norm(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def load_trials(path: str | Path) -> list[Trial]:
    """Parse the trial spreadsheet into Trial records.

    Tolerates header reordering and whitespace. Rows whose NCT cell is empty or
    malformed are still returned, with the problem recorded in ``notes``, so the
    monthly report can show what is unsearchable instead of silently dropping it.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)

    header = next(rows)
    index: dict[str, int] = {}
    for i, cell in enumerate(header):
        key = COLUMNS.get(_norm(cell))
        if key:
            index[key] = i
    if "nct" not in index:
        raise ValueError(
            f"No 'Trial Nctnumber' column found in {path}. Headers were: {header}"
        )

    trials: list[Trial] = []
    for rownum, row in enumerate(rows, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue

        def cell(key: str) -> str:
            i = index.get(key)
            return "" if i is None or i >= len(row) else str(row[i] or "").strip()

        trial = Trial(
            protocol=cell("protocol"),
            phase=cell("phase"),
            name=cell("name"),
            status=cell("status"),
            row=rownum,
        )
        raw = cell("nct").upper().replace(" ", "")
        if STRICT_NCT_RE.match(raw):
            trial.nct = raw
        elif not raw:
            trial.notes.append("no NCT number in source list")
        elif (found := NCT_RE.search(raw)):
            # e.g. "NCT12345678 / NCT87654321" or stray punctuation
            trial.nct = found.group()
            trial.notes.append(f"NCT cell was {raw!r}; used {trial.nct}")
        else:
            trial.notes.append(f"unrecognized NCT value {raw!r}")
        trials.append(trial)

    wb.close()
    return trials


def searchable(trials: list[Trial]) -> dict[str, Trial]:
    """Map NCT number -> first trial carrying it (dedupes repeat listings)."""
    out: dict[str, Trial] = {}
    for t in trials:
        nct = t.effective_nct
        if nct and nct not in out:
            out[nct] = t
    return out
