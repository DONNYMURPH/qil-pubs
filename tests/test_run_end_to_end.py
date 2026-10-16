"""Exercise the whole run path with the network mocked out."""
from pathlib import Path

import httpx
import openpyxl
import pytest
import respx

from qil_pubs.cli import main

FIXTURES = Path(__file__).parent / "fixtures"
TRIALS = FIXTURES / "trials.xlsx"

EFETCH_TEMPLATE = """<?xml version="1.0"?>
<PubmedArticleSet>
{articles}
</PubmedArticleSet>"""

ARTICLE = """ <PubmedArticle>
  <MedlineCitation>
   <PMID Version="1">{pmid}</PMID>
   <Article>
    <Journal><ISOAbbreviation>J Test</ISOAbbreviation>
      <JournalIssue><PubDate><Year>{year}</Year></PubDate></JournalIssue></Journal>
    <ArticleTitle>Paper about {nct}</ArticleTitle>
    <Abstract><AbstractText>Enrolled on {nct}.</AbstractText></Abstract>
    <AuthorList><Author><LastName>Author{pmid}</LastName><Initials>A</Initials></Author></AuthorList>
   </Article>
  </MedlineCitation>
 </PubmedArticle>"""

# NCT numbers that really are in the cancer center's list.
HITS = {"40000001": ("NCT02432235", "2024"), "40000002": ("NCT03589469", "2023")}
# Registered on ClinicalTrials.gov but never surfaced by the PubMed search.
CTGOV_ONLY = "40000003"


@pytest.fixture
def mocked_api():
    with respx.mock(assert_all_called=False) as mock:
        def esearch(request):
            # Only the batch containing our two seeded trials returns hits.
            term = request.url.params.get("term", "")
            ids = [p for p, (nct, _) in HITS.items() if nct in term]
            return httpx.Response(
                200,
                json={"esearchresult": {"count": str(len(ids)), "idlist": ids}},
            )

        def efetch(request):
            requested = request.url.params.get("id", "").split(",")
            articles = []
            for pmid in requested:
                if pmid in HITS:
                    nct, year = HITS[pmid]
                elif pmid == CTGOV_ONLY:
                    nct, year = "NCT04044768", "2022"
                else:
                    continue
                articles.append(ARTICLE.format(pmid=pmid, nct=nct, year=year))
            return httpx.Response(
                200, text=EFETCH_TEMPLATE.format(articles="\n".join(articles))
            )

        def ctgov(request):
            nct = request.url.path.rsplit("/", 1)[-1]
            refs = []
            if nct == "NCT04044768":
                refs = [
                    {"pmid": CTGOV_ONLY, "type": "RESULT", "citation": "Results paper"},
                    {"pmid": "40000099", "type": "BACKGROUND", "citation": "Prior art"},
                ]
            return httpx.Response(
                200, json={"protocolSection": {"referencesModule": {"references": refs}}}
            )

        mock.get(url__regex=r".*esearch\.fcgi.*").mock(side_effect=esearch)
        mock.get(url__regex=r".*efetch\.fcgi.*").mock(side_effect=efetch)
        mock.get(url__regex=r".*clinicaltrials\.gov/api/v2/studies/NCT.*").mock(
            side_effect=ctgov
        )
        yield mock


def _run(tmp_path, *extra):
    return main([
        "run", "--trials", str(TRIALS),
        "--out-dir", str(tmp_path / "out"),
        "--state", str(tmp_path / "out" / "state.json"),
        *extra,
    ])


def _rows(path):
    ws = openpyxl.load_workbook(path)["Publications"]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


def test_first_run_reports_everything_as_new(tmp_path, mocked_api):
    assert _run(tmp_path) == 0
    out = tmp_path / "out"
    rows = _rows(out / "qil_publications.xlsx")

    # Two from PubMed, one found only through ClinicalTrials.gov.
    assert {r["PMID"] for r in rows} == {"40000001", "40000002", CTGOV_ONLY}
    assert all(r["New this run"] == "NEW" for r in rows)

    by_pmid = {r["PMID"]: r for r in rows}
    assert by_pmid[CTGOV_ONLY]["Trial NCT(s)"] == "NCT04044768"
    assert "ctgov" in by_pmid[CTGOV_ONLY]["Matched by"]
    # A BACKGROUND citation is prior literature, not trial output.
    assert "40000099" not in {r["PMID"] for r in rows}

    # The trial name is carried over from the cancer center's own list.
    assert by_pmid["40000001"]["Trial name(s)"]
    assert (out / "state.json").exists()
    csvs = list(out.glob("new_publications_*.csv"))
    assert len(csvs) == 1
    assert csvs[0].read_text().count("\n") == 4  # header + 3 pubs


def test_second_run_reports_nothing_new(tmp_path, mocked_api):
    _run(tmp_path)
    assert _run(tmp_path) == 0
    rows = _rows(tmp_path / "out" / "qil_publications.xlsx")
    assert len(rows) == 3
    assert all(r["New this run"] in (None, "") for r in rows)
    assert all(r["First seen"] for r in rows)


def test_dry_run_leaves_state_untouched(tmp_path, mocked_api):
    _run(tmp_path, "--dry-run")
    assert not (tmp_path / "out" / "state.json").exists()
    _run(tmp_path, "--dry-run")
    rows = _rows(tmp_path / "out" / "qil_publications.xlsx")
    assert all(r["New this run"] == "NEW" for r in rows)


def test_by_trial_sheet_covers_every_searchable_trial(tmp_path, mocked_api):
    _run(tmp_path)
    wb = openpyxl.load_workbook(tmp_path / "out" / "qil_publications.xlsx")
    ws = wb["By trial"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4  # every unique NCT in the fixture, hits or not
    hits = {r[0]: r[5] for r in rows}
    assert hits["NCT02432235"] == 1
    assert hits["NCT05194293"] == 0  # searched, no publications

    attention = list(wb["Needs attention"].iter_rows(min_row=2, values_only=True))
    assert len(attention) == 1 and "no NCT" in attention[0][4]


def test_transient_pubmed_failure_is_retried(tmp_path, monkeypatch):
    monkeypatch.setattr("time.sleep", lambda *_: None)
    calls = {"n": 0}

    with respx.mock(assert_all_called=False) as mock:
        def flaky(request):
            calls["n"] += 1
            if calls["n"] == 1:
                return httpx.Response(429, text="slow down")
            return httpx.Response(200, json={"esearchresult": {"count": "0", "idlist": []}})

        mock.get(url__regex=r".*esearch\.fcgi.*").mock(side_effect=flaky)
        mock.get(url__regex=r".*efetch\.fcgi.*").mock(
            return_value=httpx.Response(200, text=EFETCH_TEMPLATE.format(articles=""))
        )
        mock.get(url__regex=r".*clinicaltrials\.gov.*").mock(
            return_value=httpx.Response(200, json={})
        )
        assert _run(tmp_path) == 0
        assert calls["n"] > 1  # the 429 was retried, not fatal
